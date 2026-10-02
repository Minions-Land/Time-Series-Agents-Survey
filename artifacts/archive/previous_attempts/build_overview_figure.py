from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
OUT='/Users/mjm/Papers/Qingsong/IJCAI2027-Time-Series-Agents-Survey/artifacts/Figure2_RSI_style.pptx'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
slide=prs.slides.add_slide(prs.slide_layouts[6])
NAVY=RGBColor(24,53,93); BLUE=RGBColor(64,112,184); PURPLE=RGBColor(111,79,157); AMBER=RGBColor(197,133,22); TEAL=RGBColor(37,126,113); GREEN=RGBColor(47,133,93); INK=RGBColor(37,45,58); MUTED=RGBColor(98,109,125); BG=RGBColor(250,251,253); PALE_BLUE=RGBColor(238,245,253); PALE_PURPLE=RGBColor(245,241,251); PALE_AMBER=RGBColor(254,248,235); PALE_GREEN=RGBColor(239,249,244)
def box(x,y,w,h,fill,line,r=MSO_SHAPE.ROUNDED_RECTANGLE,lw=1.2):
 s=slide.shapes.add_shape(r, Inches(x), Inches(y), Inches(w), Inches(h)); s.fill.solid(); s.fill.fore_color.rgb=fill; s.line.color.rgb=line; s.line.width=Pt(lw); return s
def text(x,y,w,h,txt,size=12,color=INK,bold=False,align=PP_ALIGN.LEFT,font='Aptos',margin=.06):
 tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; p=tf.paragraphs[0]; p.alignment=align; run=p.add_run(); run.text=txt; run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=color; return tb
def line(x1,y1,x2,y2,color=BLUE,lw=1.4,dash=None,arrow=False):
 c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); c.line.color.rgb=color; c.line.width=Pt(lw); 
 if dash: c.line.dash_style=dash
 if arrow: c.line.end_arrowhead=True
 return c
bg=slide.background; bg.fill.solid(); bg.fill.fore_color.rgb=BG
text(.45,.25,12.4,.34,'How a Time-Series Agent is organized',22,NAVY,True,PP_ALIGN.CENTER)
text(.45,.62,12.4,.25,'A module-level view: LLM-side capability is independent; Harness modules form a reusable stack.',10,MUTED,False,PP_ALIGN.CENTER)
box(.45,1.18,2.35,4.7,PALE_PURPLE,PURPLE,lw=1.5); text(.63,1.38,1.98,.35,'LLM side',16,PURPLE,True,PP_ALIGN.CENTER)
text(.7,1.85,1.85,.7,'Language\nunderstanding',13,PURPLE,True,PP_ALIGN.CENTER); text(.7,2.75,1.85,.7,'Temporal\nrepresentation',13,PURPLE,True,PP_ALIGN.CENTER); text(.7,3.65,1.85,.7,'Reasoning &\nexplanation',13,PURPLE,True,PP_ALIGN.CENTER); text(.64,4.75,1.98,.65,'ChatTS and related\nmodel-side works',10,MUTED,False,PP_ALIGN.CENTER); line(2.82,3.45,3.22,3.45,PURPLE,1.8,arrow=True)
box(3.2,1.08,6.35,5.0,RGBColor(255,255,255),NAVY,lw=1.8); text(3.42,1.22,5.9,.32,'Time-Series Agent runtime',16,NAVY,True,PP_ALIGN.CENTER)
box(3.48,1.72,5.8,1.05,PALE_BLUE,BLUE,lw=1.4); text(3.7,1.82,5.35,.25,'General Harness',14,BLUE,True,PP_ALIGN.CENTER); text(3.72,2.14,5.3,.25,'orchestration  ·  memory  ·  tools  ·  control  ·  verification',9,MUTED,False,PP_ALIGN.CENTER)
box(3.72,2.93,5.32,1.45,PALE_AMBER,AMBER,lw=1.4); text(3.95,3.03,4.86,.25,'General Time-Series Harness',14,AMBER,True,PP_ALIGN.CENTER)
chips=['interface','memory','tools','control','planning','optimization','coordination']
for i,ch in enumerate(chips):
 col=i%4; row=i//4; x=3.96+col*1.16; y=3.42+row*.48; w=1.05 if col<3 or row==1 else 1.12; box(x,y,w,.30,RGBColor(255,252,243),AMBER,lw=.7); text(x+.02,y+.01,w-.04,.27,ch,8,AMBER,False,PP_ALIGN.CENTER)
box(3.95,4.72,4.86,1.04,PALE_GREEN,GREEN,lw=1.4); text(4.15,4.82,4.46,.25,'Task-Specific Harness',14,GREEN,True,PP_ALIGN.CENTER); text(4.12,5.17,4.52,.34,'forecasting  ·  augmentation  ·  anomaly  ·  decision support',8.7,MUTED,False,PP_ALIGN.CENTER)
box(9.92,1.18,2.95,4.7,RGBColor(255,255,255),TEAL,lw=1.5); text(10.12,1.38,2.55,.35,'Task families',16,TEAL,True,PP_ALIGN.CENTER)
for i,(lab,sub) in enumerate([('Forecasting','predict future values'),('Augmentation','synthesize useful series'),('Anomaly','detect and diagnose'),('Decision support','act on predictions')]):
 y=1.95+i*.85; box(10.18,y,2.42,.60,PALE_GREEN,TEAL,lw=.9); text(10.28,y+.07,2.22,.22,lab,10.5,TEAL,True,PP_ALIGN.CENTER); text(10.28,y+.32,2.22,.17,sub,8,MUTED,False,PP_ALIGN.CENTER)
line(9.57,3.45,9.92,3.45,TEAL,1.8,arrow=True)
box(1.12,6.35,11.05,.72,RGBColor(255,255,255),NAVY,lw=1.3); text(1.35,6.48,1.55,.25,'Paper card',12,NAVY,True,PP_ALIGN.CENTER); text(2.95,6.46,8.95,.30,'claim  ·  evidence locator  ·  module labels  ·  task family  ·  BibTeX  ·  source status',10,MUTED,False,PP_ALIGN.CENTER)
line(6.62,2.77,6.62,2.93,BLUE,1.2,MSO_LINE_DASH_STYLE.DASH); line(6.62,4.38,6.62,4.72,AMBER,1.2,MSO_LINE_DASH_STYLE.DASH)
text(.5,7.15,12.3,.18,'Modules are recorded independently: one paper may contribute to several locations, while each contribution is assigned to its deepest applicable layer.',8.5,MUTED,False,PP_ALIGN.CENTER)
prs.save(OUT)
