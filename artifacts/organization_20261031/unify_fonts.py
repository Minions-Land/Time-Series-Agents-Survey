"""Set every text run in the given figure PPTX files to Times New Roman.

Sizes, weights, colours, and layout are unchanged; slides get an explicit
white background so that Keynote exports match the LibreOffice originals.
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor

FONT = "Times New Roman"


def shapes(collection):
    for shape in collection:
        if shape.shape_type == 6:
            yield from shapes(shape.shapes)
        else:
            yield shape


def frames(shape):
    if shape.has_text_frame:
        yield shape.text_frame
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame


for path in sys.argv[1:]:
    prs = Presentation(path)
    runs = 0
    for slide in prs.slides:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for shape in shapes(slide.shapes):
            for frame in frames(shape):
                for para in frame.paragraphs:
                    for run in para.runs:
                        run.font.name = FONT
                        runs += 1
    prs.save(path)
    print(f"{runs:4d} runs  {path}")
