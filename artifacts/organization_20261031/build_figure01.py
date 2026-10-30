"""Build Figure 1 (survey organization) as an editable one-slide PPTX.

Rows follow the title: foundations, then LLMs and harness architectures,
evaluation, and outlook. Layer colours match Figure 3 (NPG palette).
"""
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("Figure01_Organization.pptx")
FONT = "Times New Roman"
INK, MUTED, RULE = RGBColor(0x22, 0x22, 0x22), RGBColor(0x5C, 0x5C, 0x5C), RGBColor(0xC9, 0xC9, 0xC9)
GOLD, NAVY, TEAL, RED = (RGBColor.from_string(h) for h in ("E8A033", "3A5398", "00A087", "E64B35"))
SLATE = RGBColor(0x6B, 0x6B, 0x6B)
TITLE_PT, ITEM_PT, LABEL_PT = 18.5, 14.5, 14.5

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(14), Inches(6.55)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def text(box, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    for side in ("margin_left", "margin_right"):
        setattr(frame, side, Inches(0.14))
    frame.margin_top = frame.margin_bottom = Inches(0.06)
    for i, line in enumerate(runs):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.alignment = align
        for content, size, bold, color in line:
            run = para.add_run()
            run.text = content
            run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
            run.font.color.rgb = color


def section(x, y, w, h, number, title, items, accent, stacked=False):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.08
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    box.line.color.rgb, box.line.width = accent, Pt(1.25)
    box.shadow.inherit = False
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.02), Inches(y + 0.12), Inches(0.07), Inches(h - 0.24))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    bar.shadow.inherit = False
    head = [(f"§{number}  ", TITLE_PT, True, accent), (title, TITLE_PT, True, INK)]
    body = [[(item, ITEM_PT, False, MUTED)] for item in items] if stacked else [[(" · ".join(items), ITEM_PT, False, MUTED)]]
    text(box, [head, *body], anchor=MSO_ANCHOR.TOP if stacked else MSO_ANCHOR.MIDDLE)


def label(x, y, w, h, content, color=MUTED, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    text(box, [[(content, LABEL_PT, True, color)]], align=align)


def arrow(x1, y1, x2, y2, color=SLATE):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb, line.line.width = color, Pt(1.5)
    ln = line.line._get_or_add_ln()
    ln.append(etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", type="triangle", w="med", len="med"))


def rule(y):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.3), Inches(y), Inches(13.7), Inches(y))
    line.line.color.rgb, line.line.width = RULE, Pt(0.75)


LEFT, RIGHT = 1.75, 13.7
label(0.3, 0.32, 1.45, 0.9, "Foundations")
section(LEFT, 0.3, 5.3, 0.95, 1, "Introduction", ["construction question", "taxonomy", "contributions"], SLATE)
section(LEFT + 5.5, 0.3, RIGHT - LEFT - 5.5, 0.95, 2, "What Is a Time Series Agent?", ["definition", "scope of inclusion", "layer assignment"], SLATE)
rule(1.45)

y2, h2 = 2.0, 1.95
label(0.3, y2 + 0.5, 1.45, 0.9, "Components")
gap = 0.38
widths = [2.7, 2.95, 2.65]
xs = [LEFT, LEFT + widths[0] + 0.3]
xs.append(xs[1] + widths[1] + gap)
xs.append(xs[2] + widths[2] + gap)
widths.append(RIGHT - xs[3])
label(LEFT - 0.14, 1.55, widths[0], 0.4, "LLMs")
label(xs[1] - 0.14, 1.55, RIGHT - xs[1], 0.4, "Harness architectures, from general to specific")
specs = [(3, "LLM-Side", ["temporal representation", "instruction and reasoning", "workflow learning"], GOLD),
         (4, "General Harness", ["interface and interaction", "memory and context", "tools and execution", "planning and orchestration", "optimization and evolution"], NAVY),
         (5, "General Time Series Harness", ["data management", "temporal processing", "context and memory", "temporal validation"], TEAL),
         (6, "Task-Specific Harness", ["forecasting", "synthesis", "anomaly diagnosis", "decision support"], RED)]
for x, w, (number, title, items, accent) in zip(xs, widths, specs):
    section(x, y2, w, h2, number, title, items, accent, stacked=True)
for k in (1, 2):
    xa = xs[k] + widths[k]
    arrow(xa + 0.04, y2 + h2 / 2, xa + gap - 0.04, y2 + h2 / 2)
rule(4.15)

label(0.3, 4.3, 1.45, 0.9, "Evaluation")
section(LEFT, 4.28, RIGHT - LEFT, 0.85, 7, "Infrastructure and Benchmarks",
        ["evaluation contract", "benchmark landscape", "temporal question answering", "reproducible infrastructure", "benchmark selection"], SLATE)
rule(5.33)

label(0.3, 5.5, 1.45, 0.9, "Outlook")
section(LEFT, 5.48, 6.6, 0.85, 8, "Open Problems and System Analyses", ["reusable temporal infrastructure", "open problems", "design implications"], SLATE)
section(LEFT + 6.8, 5.48, RIGHT - LEFT - 6.8, 0.85, 9, "Conclusion", ["design judgments", "alignment of the three parts"], SLATE)

prs.save(OUT)
print(OUT)
